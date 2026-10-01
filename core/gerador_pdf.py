"""
Gera o Excel do pedido (baseado no molde) e exporta pra PDF via Excel COM.
"""
import shutil
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Border, Side, Alignment, Font, PatternFill
from openpyxl.worksheet.properties import PageSetupProperties

from core.config import RAIZ


PASTA_TEMPLATES = RAIZ / "templates"
MOLDE_PEDIDO = PASTA_TEMPLATES / "modelo_pedido.xlsx"


def _formata_moeda(v) -> str:
    s = f"{v:,.2f}"
    s = s.replace(",", "X").replace(".", ",").replace("X", ".")
    return f"R${s}"


def gerar_excel_pedido(itens: list, fornecedor_codigo: str,
                        fornecedor_nome: str, filial_numero: int,
                        config: dict, caminho_destino: Path):
    if not MOLDE_PEDIDO.exists():
        raise FileNotFoundError(f"Molde não encontrado: {MOLDE_PEDIDO}")

    shutil.copy2(MOLDE_PEDIDO, caminho_destino)

    wb = openpyxl.load_workbook(caminho_destino)
    ws = wb.active

    # 1. Desfaz mesclagens das linhas >= 9
    merges_para_remover = []
    for merge in list(ws.merged_cells.ranges):
        if merge.min_row >= 9:
            merges_para_remover.append(str(merge))
    for m in merges_para_remover:
        ws.unmerge_cells(m)

    # 2. Limpa da linha 9 em diante
    max_row = ws.max_row
    for r in range(9, max_row + 1):
        for c in range(1, 11):
            ws.cell(row=r, column=c).value = None

    # 3. Substitui nome / CNPJ / data / fornecedor
    filial_info = config.get("filiais", {}).get(str(filial_numero), {})
    nome_filial = filial_info.get("nome", f"FILIAL {filial_numero}")
    cnpj = filial_info.get("cnpj", "")

    ws["B2"] = nome_filial
    ws["B3"] = None
    ws["B4"] = f"CNPJ: {cnpj}"
    ws["I3"] = date.today().strftime("%d/%m/%Y")

    try:
        ws["G3"] = int(fornecedor_codigo)
    except (ValueError, TypeError):
        ws["G3"] = fornecedor_codigo

    # Nome do fornecedor (G5, mesclado com H5:I5)
    ws["G5"] = fornecedor_nome
    ws["G5"].font = Font(bold=True)
    ws["G5"].alignment = Alignment(horizontal="center")

    # 4. Preenche os produtos
    from copy import copy

    linha_inicial = 9
    linha_modelo = 9  # linha 9 do molde tem a formatação base

    for i, it in enumerate(itens):
        r = linha_inicial + i

        # Copia estilo da linha 9 do molde pra linha atual (se não for a própria linha 9)
        if r != linha_modelo:
            for col in range(1, 10):
                origem = ws.cell(row=linha_modelo, column=col)
                destino = ws.cell(row=r, column=col)
                if origem.has_style:
                    destino._style = copy(origem._style)

        ws.cell(row=r, column=1, value=it.get("codigo_barras", ""))
        ws.cell(row=r, column=2, value=it.get("codigo_interno", ""))
        ws.cell(row=r, column=3, value=it.get("codigo_no_fornecedor", ""))
        ws.cell(row=r, column=4, value=it.get("descricao", ""))
        ws.cell(row=r, column=7, value=int(it["qtd"]))
        cel_preco = ws.cell(row=r, column=8, value=float(it["preco"]))
        cel_preco.number_format = '"R$"        #,##0.00'
        cel_preco.alignment = Alignment(horizontal="right")

        cel_total = ws.cell(row=r, column=9, value=float(it["total"]))
        cel_total.number_format = '"R$"        #,##0.00'
        cel_total.alignment = Alignment(horizontal="right")

        ws.cell(row=r, column=7).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=8).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=9).alignment = Alignment(horizontal="center")

        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=6)

    # 5. Linha de totais
    linha_totais = linha_inicial + len(itens)

    ws.cell(row=linha_totais, column=1, value="QTD DE ITENS")
    ws.cell(row=linha_totais, column=1).font = Font(bold=True)
    ws.cell(row=linha_totais, column=1).alignment = Alignment(horizontal="center")

    ws.cell(row=linha_totais, column=3, value=len(itens))
    ws.cell(row=linha_totais, column=3).font = Font(bold=True)
    ws.cell(row=linha_totais, column=3).alignment = Alignment(horizontal="center")

    ws.cell(row=linha_totais, column=7, value="VALOR TOTAL")
    ws.cell(row=linha_totais, column=7).font = Font(bold=True)
    ws.cell(row=linha_totais, column=7).alignment = Alignment(horizontal="center")

    

    fill_total = PatternFill("solid", fgColor="DCE6F1")
    for col in range(1, 10):
        ws.cell(row=linha_totais, column=col).fill = fill_total

    ws.merge_cells(start_row=linha_totais, start_column=1,
                    end_row=linha_totais, end_column=2)
    ws.merge_cells(start_row=linha_totais, start_column=7,
                    end_row=linha_totais, end_column=8)

    # 6. Larguras
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 40
    ws.column_dimensions["E"].width = 2
    ws.column_dimensions["F"].width = 2
    ws.column_dimensions["G"].width = 10
    ws.column_dimensions["H"].width = 12
    ws.column_dimensions["I"].width = 14

    # 7. Área de impressão + fit
    ws.print_area = f"A1:I{linha_totais}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

    wb.save(caminho_destino)
    return caminho_destino


def exportar_pdf(caminho_excel: Path, caminho_pdf: Path):
    import win32com.client
    import pythoncom

    pythoncom.CoInitialize()

    excel = win32com.client.Dispatch("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False

    try:
        wb = excel.Workbooks.Open(str(caminho_excel.resolve()))
        wb.ExportAsFixedFormat(
            Type=0,
            Filename=str(caminho_pdf.resolve()),
            Quality=0,
            IncludeDocProperties=True,
            IgnorePrintAreas=False,
            OpenAfterPublish=False,
        )
        wb.Close(SaveChanges=False)
    finally:
        excel.Quit()
        pythoncom.CoUninitialize()

    return caminho_pdf