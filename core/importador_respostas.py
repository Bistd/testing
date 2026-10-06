"""
Importador BLINDADO do Excel de resposta dos fornecedores.
Aceita variações de formato, nome de arquivo, cabeçalho, etc.
"""
from pathlib import Path
import openpyxl


def ler_resposta_fornecedor(caminho: Path) -> dict:
    """
    Lê um Excel devolvido por fornecedor.

    Retorna:
        {
            "codigo_fornecedor_c3": "4" ou None,   # lido da célula C3
            "itens": [
                {"codigo_barras": "...", "descricao": "...", "preco": 12.50},
                ...
            ],
            "erros": [...],
            "total_itens": N,
        }
    """
    resultado = {
        "codigo_fornecedor_c3": None,
        "itens": [],
        "erros": [],
        "total_itens": 0,
    }

    try:
        wb = openpyxl.load_workbook(str(caminho), data_only=True)
    except Exception as e:
        resultado["erros"].append(f"Não consegui abrir o arquivo: {e}")
        return resultado

    # Pega a primeira aba ativa
    ws = wb.active

    # ============================================
    # 1. Tenta ler o código do fornecedor da célula C3
    # ============================================
    try:
        valor_c3 = ws["C3"].value
        if valor_c3 is not None:
            # Normaliza (pode vir como float 4.0)
            if isinstance(valor_c3, float) and valor_c3.is_integer():
                valor_c3 = int(valor_c3)
            resultado["codigo_fornecedor_c3"] = str(valor_c3).strip()
    except Exception:
        pass

    # ============================================
    # 2. Procura os produtos no arquivo
    # ============================================
    # Estratégia: varre TODAS as linhas procurando o padrão
    # "código | descrição | preço"
    # Aceita linhas em qualquer posição

    for row in ws.iter_rows():
        try:
            # Pega as 3 primeiras colunas
            col_a = row[0].value if len(row) > 0 else None
            col_b = row[1].value if len(row) > 1 else None
            col_c = row[2].value if len(row) > 2 else None

            if col_a is None or col_b is None:
                continue

            # Ignora cabeçalhos
            if _e_cabecalho(col_a) or _e_cabecalho(col_b):
                continue

            # Ignora linhas de bloco (nome da filial, fornecedores, etc)
            if _e_linha_de_bloco(col_b):
                continue

            # Se a coluna C tem preço, é um item válido
            preco = _extrair_preco(col_c)
            if preco is None or preco <= 0:
                continue

            codigo_barras = _limpar_texto(col_a)
            descricao = _limpar_texto(col_b)

            resultado["itens"].append({
                "codigo_barras": codigo_barras,
                "descricao": descricao,
                "preco": preco,
            })

        except Exception as e:
            resultado["erros"].append(f"Erro numa linha: {e}")
            continue

    resultado["total_itens"] = len(resultado["itens"])
    wb.close()
    return resultado


# ============================================================
# HELPERS
# ============================================================

def _e_cabecalho(valor) -> bool:
    """Verifica se o valor é um cabeçalho (não é código nem descrição)."""
    if valor is None:
        return False
    s = str(valor).strip().upper()
    cabecalhos = {
        "COD", "CÓD", "CODIGO", "CÓDIGO", "COD BARRAS", "CÓD BARRAS",
        "COD/BARRAS", "CÓD/BARRAS", "PRODUTO", "DESCRICAO", "DESCRIÇÃO",
        "VL UNIT", "VL UNIT.", "PRECO", "PREÇO", "VALOR", "VL",
    }
    return s in cabecalhos


def _e_linha_de_bloco(valor) -> bool:
    """Verifica se é linha de cabeçalho de bloco (nome de filial, fornecedores)."""
    if valor is None:
        return False
    s = str(valor).strip().upper()
    # Nomes de filial (do config) ou cabeçalhos especiais
    return (
        "COMERCIAL" in s or "SUPERMERCADO" in s or
        "FILIAL" in s or "FORNECEDOR" in s
    )


def _extrair_preco(valor):
    """Tenta converter um valor em preço (float)."""
    if valor is None:
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    s = str(valor).strip()
    if not s:
        return None
    # Remove R$, pontos de milhar, troca vírgula por ponto
    s = s.replace("R$", "").replace(" ", "").strip()
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _limpar_texto(valor) -> str:
    """Converte valor em string limpa (remove .0 de floats)."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()
