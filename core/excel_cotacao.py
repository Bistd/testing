"""
Gera e lê os Excels de cotação (trabalho + cego).
"""
from pathlib import Path
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

from core.config import RAIZ
from core.util_meses import ultimos_n_meses, MESES_NUMERO


PASTA_TEMPLATES = RAIZ / "templates"
PASTA_TRABALHO = RAIZ / "excel_trabalho"
PASTA_BACKUP = PASTA_TRABALHO / "backups"


def _garantir_pastas():
    PASTA_TEMPLATES.mkdir(parents=True, exist_ok=True)
    PASTA_TRABALHO.mkdir(parents=True, exist_ok=True)
    PASTA_BACKUP.mkdir(parents=True, exist_ok=True)


# ============================================================
# ESTILOS
# ============================================================

COR_VERMELHO = "C00000"
COR_CINZA = "F2F2F2"
COR_AZUL = "1F4E79"
COR_AMARELO_QTD = "FFF2CC"     # QTD
COR_AMARELO_CX = "FCE4B6"      # CX (mais alaranjado)
COR_AMARELO_EST = "FAE3B8"     # ESTOQUE (parecido mas diferente)
COR_BRANCO = "FFFFFF"

FONTE_TITULO = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
FONTE_CABECALHO = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
FONTE_NORMAL = Font(name="Calibri", size=10)
FONTE_COD = Font(name="Consolas", size=9)
FONTE_BOLD = Font(name="Calibri", size=10, bold=True)

FILL_VERMELHO = PatternFill("solid", fgColor=COR_VERMELHO)
FILL_CINZA = PatternFill("solid", fgColor=COR_CINZA)
FILL_AZUL = PatternFill("solid", fgColor=COR_AZUL)
FILL_AMARELO_QTD = PatternFill("solid", fgColor=COR_AMARELO_QTD)
FILL_AMARELO_CX = PatternFill("solid", fgColor=COR_AMARELO_CX)
FILL_AMARELO_EST = PatternFill("solid", fgColor=COR_AMARELO_EST)
FILL_BRANCO = PatternFill("solid", fgColor=COR_BRANCO)

BORDA_FINA = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)

BORDA_GRUPO = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="medium", color="333333"),
)

CENTRO = Alignment(horizontal="center", vertical="center")
ESQUERDA = Alignment(horizontal="left", vertical="center")
DIREITA = Alignment(horizontal="right", vertical="center")


# ============================================================
# EXCEL DE TRABALHO
# ============================================================

def gerar_excel_trabalho(produtos: list, ano_mes_atual: str,
                          caminho_destino: Path, config: dict):
    """
    NOVA ORDEM:
      A=1 COD | B=2 PRODUTO | C=3 E/S/A | D..O=4..15 meses (12)
      P=16 CX | Q=17 ESTOQUE | R=18 QTD (mesclada) | S=19 VL UNIT (mesclada)
      U=21 auxiliar (escondida)
    """
    _garantir_pastas()

    filial_id = config.get("filial_ativa", 1)
    nome_filial = config.get("filiais", {}).get(str(filial_id), {}).get("nome", f"FILIAL {filial_id}")

    lista_meses = ultimos_n_meses(ano_mes_atual, 12)
    siglas = [MESES_NUMERO[int(m.split("-")[1])] for m in lista_meses]
    anos = [m.split("-")[0] for m in lista_meses]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cotação"

    # ---- Larguras ----
    larguras = {
        "A": 12, "B": 45, "C": 6,
        "P": 6, "Q": 10, "R": 12, "S": 12,
        "U": 12,
    }
    for col in ["D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N", "O"]:
        larguras[col] = 8
    for col, w in larguras.items():
        ws.column_dimensions[col].width = w

    # ---- LINHA 1: Título ----
    ws.merge_cells("A1:S1")
    c = ws["A1"]
    c.value = f"COTAÇÃO — {nome_filial}"
    c.font = FONTE_TITULO
    c.fill = FILL_VERMELHO
    c.alignment = CENTRO
    ws.row_dimensions[1].height = 28

    # ---- LINHA 2: Subtítulo ----
    ws.merge_cells("A2:E2")
    c = ws["A2"]
    c.value = f"Mês: {siglas[-1]}/{anos[-1]}"
    c.font = Font(name="Calibri", size=11, bold=True, color="333333")
    c.alignment = ESQUERDA

    ws.merge_cells("F2:H2")
    c = ws["F2"]
    c.value = "Total de produtos:"
    c.font = FONTE_BOLD
    c.alignment = DIREITA

    c = ws["I2"]
    c.value = len(produtos)
    c.font = FONTE_BOLD
    c.alignment = ESQUERDA

    ws.merge_cells("J2:M2")
    c = ws["J2"]
    c.value = "Valor total estimado:"
    c.font = FONTE_BOLD
    c.alignment = DIREITA

    ws.merge_cells("N2:S2")
    c = ws["N2"]
    c.value = "=IFERROR(SUMPRODUCT(U3:U100000),0)"
    c.font = Font(name="Calibri", size=11, bold=True, color=COR_VERMELHO)
    c.number_format = 'R$ #,##0.00'
    c.alignment = ESQUERDA

    ws.row_dimensions[2].height = 22

    # ---- LINHAS 3-4: Cabeçalho ----
    ws.merge_cells("A3:A4"); ws["A3"] = "COD"
    ws.merge_cells("B3:B4"); ws["B3"] = "PRODUTO"
    ws.merge_cells("C3:C4"); ws["C3"] = "E/S/A"

    # Meses agrupados por ano
    ano_atual = None
    inicio_mes = None
    for i, (m, a) in enumerate(zip(lista_meses, anos)):
        col_idx = 4 + i
        if a != ano_atual:
            if ano_atual is not None:
                ws.merge_cells(start_row=3, start_column=inicio_mes,
                               end_row=3, end_column=col_idx - 1)
                cc = ws.cell(row=3, column=inicio_mes)
                cc.value = f"{ano_atual}"
                cc.font = FONTE_CABECALHO
                cc.fill = FILL_VERMELHO
                cc.alignment = CENTRO
            ano_atual = a
            inicio_mes = col_idx

    if inicio_mes is not None:
        ws.merge_cells(start_row=3, start_column=inicio_mes,
                       end_row=3, end_column=15)
        cc = ws.cell(row=3, column=inicio_mes)
        cc.value = f"{ano_atual}"
        cc.font = FONTE_CABECALHO
        cc.fill = FILL_VERMELHO
        cc.alignment = CENTRO

    # Cabeçalhos fixos (NOVA ORDEM)
    ws.merge_cells("P3:P4"); ws["P3"] = "CX"
    ws.merge_cells("Q3:Q4"); ws["Q3"] = "ESTOQUE"
    ws.merge_cells("R3:R4"); ws["R3"] = "QTD"
    ws.merge_cells("S3:S4"); ws["S3"] = "VL UNIT"

    # Estilos no cabeçalho
    for col in range(1, 20):
        for linha in [3, 4]:
            cel = ws.cell(row=linha, column=col)
            cel.fill = FILL_VERMELHO
            cel.font = FONTE_CABECALHO
            cel.alignment = CENTRO
            cel.border = BORDA_FINA

    for i, s in enumerate(siglas):
        ws.cell(row=4, column=4 + i, value=s)

    ws.row_dimensions[3].height = 18
    ws.row_dimensions[4].height = 20

    # ---- DADOS ----
    linha = 5
    for p in produtos:
        hist = p.get("hist", {})
        linha_e = linha
        linha_s = linha + 1
        linha_a = linha + 2
        linha_prox = linha + 3

        # ==================== LINHA E ====================
        ws.cell(row=linha_e, column=1, value=p.get("codigo_interno", "")).font = FONTE_COD
        ws.cell(row=linha_e, column=2, value=p.get("descricao", "")).font = FONTE_NORMAL
        ws.cell(row=linha_e, column=3, value="E").font = FONTE_BOLD
        ws.cell(row=linha_e, column=3).alignment = CENTRO

        for j, m in enumerate(lista_meses):
            v = hist.get(m, {}).get("qtd_entrada", 0) if hist else 0
            c = ws.cell(row=linha_e, column=4 + j, value=v)
            c.font = FONTE_NORMAL
            c.alignment = CENTRO

        # ==================== LINHA S ====================
        ws.cell(row=linha_s, column=1, value=p.get("codigo_barras", "")).font = FONTE_COD
        ws.cell(row=linha_s, column=3, value="S").font = FONTE_BOLD
        ws.cell(row=linha_s, column=3).alignment = CENTRO

        for j, m in enumerate(lista_meses):
            v = hist.get(m, {}).get("qtd_venda", 0) if hist else 0
            c = ws.cell(row=linha_s, column=4 + j, value=v)
            c.font = FONTE_NORMAL
            c.alignment = CENTRO

        # ==================== LINHA A (Avaria) ====================
        ws.cell(row=linha_a, column=3, value="A").font = FONTE_BOLD
        ws.cell(row=linha_a, column=3).alignment = CENTRO

        for j, m in enumerate(lista_meses):
            v = hist.get(m, {}).get("qtd_avaria", 0) if hist else 0
            c = ws.cell(row=linha_a, column=4 + j, value=v)
            c.font = FONTE_NORMAL
            c.alignment = CENTRO

        # ==================== COLUNAS MESCLADAS (E + S + A) ====================
        # CX (P=16)
        ws.merge_cells(start_row=linha_e, start_column=16,
                       end_row=linha_a, end_column=16)
        cel_cx = ws.cell(row=linha_e, column=16, value=p.get("cx", 1))
        cel_cx.number_format = '0'
        cel_cx.fill = FILL_AMARELO_CX
        cel_cx.alignment = CENTRO
        cel_cx.font = FONTE_BOLD

        # ESTOQUE (Q=17)
        ws.merge_cells(start_row=linha_e, start_column=17,
                       end_row=linha_a, end_column=17)
        cel_est = ws.cell(row=linha_e, column=17, value=int(p.get("estoque", 0)))
        cel_est.number_format = '0'
        cel_est.fill = FILL_AMARELO_EST
        cel_est.alignment = CENTRO
        cel_est.font = FONTE_BOLD

        # QTD (R=18) — editable
        ws.merge_cells(start_row=linha_e, start_column=18,
                       end_row=linha_a, end_column=18)
        cel_qtd = ws.cell(row=linha_e, column=18, value=0)
        cel_qtd.number_format = '0'
        cel_qtd.fill = FILL_AMARELO_QTD
        cel_qtd.alignment = CENTRO
        cel_qtd.font = FONTE_BOLD

        # VL UNIT (S=19)
        ws.merge_cells(start_row=linha_e, start_column=19,
                       end_row=linha_a, end_column=19)
        cel_vl = ws.cell(row=linha_e, column=19, value=p.get("vl_unit", 0))
        cel_vl.number_format = 'R$ #,##0.00'
        cel_vl.alignment = CENTRO
        cel_vl.font = FONTE_NORMAL

        # AUXILIAR (U=21) — escondida
        cel_aux = ws.cell(row=linha_e, column=21)
        cel_aux.value = f"=IF(R{linha_e}=\"\",0,R{linha_e})*S{linha_e}"
        cel_aux.number_format = 'R$ #,##0.00'

        # ==================== ESTILOS E BORDAS ====================
        for col in range(1, 20):
            if col in (16, 17, 18, 19):
                continue  # colunas mescladas tratadas depois

            for ln in [linha_e, linha_s, linha_a]:
                cel = ws.cell(row=ln, column=col)

                # Borda de grupo (última linha do produto)
                if ln == linha_a:
                    cel.border = BORDA_GRUPO
                else:
                    cel.border = BORDA_FINA

                # Cor de fundo alternada
                if ln == linha_e:
                    cel.fill = FILL_CINZA
                elif ln == linha_s:
                    cel.fill = FILL_BRANCO
                else:  # linha A
                    cel.fill = PatternFill("solid", fgColor="FEF9F3")

        # Borda de grupo nas colunas mescladas
        for col_m in (16, 17, 18, 19):
            ws.cell(row=linha_a, column=col_m).border = BORDA_GRUPO
            ws.cell(row=linha_e, column=col_m).border = BORDA_FINA

        linha = linha_prox  # pula 3 linhas (E, S, A)

    # ---- Esconde coluna auxiliar ----
    ws.column_dimensions["U"].hidden = True

    # ---- Congela painéis ----
    ws.freeze_panes = "D5"

    # ---- Ajusta largura da descrição ----
    _ajustar_largura_descricao(ws, col_idx=2, min_width=25, max_width=60)

    wb.save(caminho_destino)
    return caminho_destino


def _ajustar_largura_descricao(ws, col_idx=2, min_width=25, max_width=60):
    maior = min_width
    for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
        for cell in row:
            if cell.value:
                t = len(str(cell.value))
                if t > maior:
                    maior = t
    ws.column_dimensions[get_column_letter(col_idx)].width = min(maior + 3, max_width)


# ============================================================
# LEITURA — Excel de trabalho (QTD está em R=18 agora)
# ============================================================

def ler_excel_trabalho(caminho: Path) -> list:
    """
    Lê a QTD (coluna R=18) das linhas E.
    """
    wb = openpyxl.load_workbook(caminho, data_only=True)
    ws = wb.active

    resultado = []
    for r in range(5, ws.max_row + 1):
        marcador = ws.cell(row=r, column=3).value
        if marcador != "E":
            continue

        codigo = ws.cell(row=r, column=1).value
        qtd = ws.cell(row=r, column=18).value  # R = 18

        if codigo is None:
            continue

        try:
            codigo_str = str(int(codigo)) if isinstance(codigo, float) else str(codigo).strip()
        except (ValueError, TypeError):
            codigo_str = str(codigo).strip()

        try:
            qtd_val = float(qtd) if qtd is not None else 0.0
        except (ValueError, TypeError):
            qtd_val = 0.0

        resultado.append({"codigo_interno": codigo_str, "qtd": qtd_val})

    wb.close()
    return resultado


# ============================================================
# BACKUP
# ============================================================

def fazer_backup_excel(caminho: Path, rotulo: str = "") -> Path:
    _garantir_pastas()
    if not caminho.exists():
        return None
    ts = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    nome = f"{caminho.stem}_{rotulo}_{ts}{caminho.suffix}"
    destino = PASTA_BACKUP / nome
    import shutil
    shutil.copy2(caminho, destino)
    backups = sorted(PASTA_BACKUP.glob("*.xlsx"), key=lambda p: p.stat().st_mtime)
    while len(backups) > 20:
        backups.pop(0).unlink()
    return destino


# ============================================================
# HELPERS
# ============================================================

def listar_fornecedores_do_banco() -> list:
    """Retorna fornecedores QUE PARTICIPAM DA COTAÇÃO."""
    from core.database import conectar
    conn = conectar()
    cur = conn.cursor()
    cur.execute("""
        SELECT codigo, nome FROM fornecedores
        WHERE participa_cotacao = 1 AND ativo = 1
        ORDER BY CAST(codigo AS INTEGER), codigo
    """)
    resultado = [{"codigo": r["codigo"], "nome": r["nome"] or ""} for r in cur.fetchall()]
    conn.close()
    return resultado


# ============================================================
# EXCEL CEGO (mantém igual, mas usa lista filtrada de fornecedores)
# ============================================================

def gerar_excel_cego(produtos_f1: list, produtos_f2: list,
                      fornecedores: list, config: dict, caminho_destino: Path):
    """
    Gera o Excel cego (o que vai pros fornecedores).
    Só inclui a lista de fornecedores QUE PARTICIPAM DA COTAÇÃO.
    """
    _garantir_pastas()

    nome_f1 = config.get("filiais", {}).get("1", {}).get("nome", "FILIAL 1")
    nome_f2 = config.get("filiais", {}).get("2", {}).get("nome", "FILIAL 2")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cotação"

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 55
    ws.column_dimensions["C"].width = 15
    ws.column_dimensions["D"].width = 3
    ws.column_dimensions["E"].width = 8
    ws.column_dimensions["F"].width = 38

    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value = "COTAÇÃO BISTECÃO"
    c.font = FONTE_TITULO
    c.fill = FILL_VERMELHO
    c.alignment = CENTRO
    ws.row_dimensions[1].height = 28

    ws["A3"] = "INSERIR CÓDIGO DO FORNECEDOR:"
    ws["A3"].font = FONTE_BOLD
    cel_cod = ws["C3"]
    cel_cod.fill = PatternFill("solid", fgColor="FFFF00")
    cel_cod.font = Font(bold=True, size=14, color="000000")
    cel_cod.alignment = CENTRO
    cel_cod.number_format = '0'
    cel_cod.border = Border(
        left=Side(style="medium"), right=Side(style="medium"),
        top=Side(style="medium"), bottom=Side(style="medium"),
    )

    ws["A4"] = "FORNECEDOR:"
    ws["A4"].font = Font(bold=True, size=11, color="FFFFFF")
    ws["A4"].fill = FILL_AZUL
    ws["A4"].alignment = ESQUERDA

    ws.merge_cells("B4:F4")
    cel_nome = ws["B4"]
    cel_nome.value = (
        '=IF(C3="", "⚠ PREENCHA O CÓDIGO DO FORNECEDOR",'
        ' IFERROR(VLOOKUP(C3, $E$6:$F$1000, 2, FALSE),'
        ' "⚠ CÓDIGO NÃO ENCONTRADO"))'
    )
    cel_nome.font = Font(bold=True, size=11, color="FFFFFF")
    cel_nome.fill = FILL_AZUL
    cel_nome.alignment = CENTRO

    ws.row_dimensions[3].height = 24
    ws.row_dimensions[4].height = 22

    # Lista lateral
    linha_forn = 5
    ws.merge_cells(start_row=linha_forn, start_column=5,
                   end_row=linha_forn, end_column=6)
    c = ws.cell(row=linha_forn, column=5, value="FORNECEDORES DA COTAÇÃO")
    c.font = FONTE_CABECALHO
    c.fill = FILL_AZUL
    c.alignment = CENTRO
    linha_forn += 1

    c = ws.cell(row=linha_forn, column=5, value="CÓD")
    c.font = FONTE_CABECALHO
    c.fill = FILL_VERMELHO
    c.alignment = CENTRO
    c.border = BORDA_FINA

    c = ws.cell(row=linha_forn, column=6, value="FORNECEDOR")
    c.font = FONTE_CABECALHO
    c.fill = FILL_VERMELHO
    c.alignment = CENTRO
    c.border = BORDA_FINA
    linha_forn += 1

    for f in fornecedores:
        codigo = f.get("codigo", "")
        try:
            codigo = int(codigo)
        except (ValueError, TypeError):
            pass
        cel = ws.cell(row=linha_forn, column=5, value=codigo)
        cel.font = FONTE_NORMAL
        cel.alignment = CENTRO
        cel.number_format = '0'
        cel.border = BORDA_FINA
        ws.cell(row=linha_forn, column=6, value=f.get("nome", "")).font = FONTE_NORMAL
        ws.cell(row=linha_forn, column=6).border = BORDA_FINA
        linha_forn += 1

    # Bloco F1
    linha = 7
    ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=3)
    c = ws.cell(row=linha, column=1, value=nome_f1)
    c.font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    c.fill = FILL_AZUL
    c.alignment = CENTRO
    ws.row_dimensions[linha].height = 22
    linha += 1

    for i, h in enumerate(["COD BARRAS", "PRODUTO", "VL UNIT"], start=1):
        c = ws.cell(row=linha, column=i, value=h)
        c.font = FONTE_CABECALHO
        c.fill = FILL_VERMELHO
        c.alignment = CENTRO
        c.border = BORDA_FINA
    linha += 1

    primeira_linha_f1 = linha
    for p in produtos_f1:
        ws.cell(row=linha, column=1, value=p.get("codigo_barras", "")).font = FONTE_COD
        ws.cell(row=linha, column=2, value=p.get("descricao", "")).font = FONTE_NORMAL
        cel = ws.cell(row=linha, column=3, value=None)
        cel.number_format = 'R$ #,##0.00'
        cel.alignment = CENTRO
        for col in range(1, 4):
            ws.cell(row=linha, column=col).border = BORDA_FINA
        linha += 1
    ultima_linha_f1 = linha - 1
    linha += 1

    # Bloco F2
    ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=3)
    c = ws.cell(row=linha, column=1, value=nome_f2)
    c.font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    c.fill = FILL_AZUL
    c.alignment = CENTRO
    ws.row_dimensions[linha].height = 22
    linha += 1

    for i, h in enumerate(["COD BARRAS", "PRODUTO", "VL UNIT"], start=1):
        c = ws.cell(row=linha, column=i, value=h)
        c.font = FONTE_CABECALHO
        c.fill = FILL_VERMELHO
        c.alignment = CENTRO
        c.border = BORDA_FINA
    linha += 1

    primeira_linha_f2 = linha
    for p in produtos_f2:
        ws.cell(row=linha, column=1, value=p.get("codigo_barras", "")).font = FONTE_COD
        ws.cell(row=linha, column=2, value=p.get("descricao", "")).font = FONTE_NORMAL
        cel = ws.cell(row=linha, column=3, value=None)
        cel.number_format = 'R$ #,##0.00'
        cel.alignment = CENTRO
        for col in range(1, 4):
            ws.cell(row=linha, column=col).border = BORDA_FINA
        linha += 1
    ultima_linha_f2 = linha - 1

    # Formatação condicional
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.worksheet.datavalidation import DataValidation

    if primeira_linha_f1 <= ultima_linha_f2:
        range_precos = f"C{primeira_linha_f1}:C{ultima_linha_f2}"
        rule_cinza = FormulaRule(
            formula=['$C$3=""'],
            fill=PatternFill("solid", fgColor="E0E0E0"),
            font=Font(color="999999"),
            stopIfTrue=False,
        )
        ws.conditional_formatting.add(range_precos, rule_cinza)

        rule_amarela = FormulaRule(
            formula=['$C$3<>""'],
            fill=PatternFill("solid", fgColor="FFF2CC"),
            stopIfTrue=False,
        )
        ws.conditional_formatting.add(range_precos, rule_amarela)

        dv = DataValidation(
            type="custom",
            formula1='=$C$3<>""',
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="⚠ Código do fornecedor",
            error="Preencha o código do fornecedor antes de digitar o preço!",
        )
        ws.add_data_validation(dv)
        dv.add(range_precos)

        rule_nome_vermelho = FormulaRule(
            formula=['OR($C$3="", ISERROR(VLOOKUP($C$3,$E$6:$F$1000,2,FALSE)))'],
            fill=PatternFill("solid", fgColor="C00000"),
            font=Font(bold=True, color="FFFFFF"),
        )
        ws.conditional_formatting.add("B4", rule_nome_vermelho)

        rule_nome_verde = FormulaRule(
            formula=['AND($C$3<>"", NOT(ISERROR(VLOOKUP($C$3,$E$6:$F$1000,2,FALSE))))'],
            fill=PatternFill("solid", fgColor="2E7D32"),
            font=Font(bold=True, color="FFFFFF"),
        )
        ws.conditional_formatting.add("B4", rule_nome_verde)

    ws.freeze_panes = "A5"

    # Força o Excel a recalcular tudo ao abrir
    wb.calculation = CalcProperties(fullCalcOnLoad=True)

    wb.save(caminho_destino)
    return caminho_destino